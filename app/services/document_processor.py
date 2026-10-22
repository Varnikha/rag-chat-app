from io import BytesIO
from pathlib import Path
from typing import List, Tuple, Union

from PyPDF2 import PdfReader
from docx import Document as DocxDocument
from sqlalchemy.orm import Session

from app.models.user import Document as UserDocument
from app.services.text_chunker import TextChunker


class DocumentProcessor:
    """Process uploaded documents and extract text for RAG indexing."""

    def __init__(self):
        self.chunker = TextChunker()

    def extract_text(self, file_source: Union[str, Path, bytes, BytesIO]) -> str:
        """Read and extract text from a file path or in-memory file-like object."""
        if isinstance(file_source, bytes):
            return self._extract_text_from_bytes(file_source, "unknown")

        if hasattr(file_source, "read"):
            return self._extract_text_from_bytes(file_source.read(), getattr(file_source, "name", "unknown"))

        path = Path(file_source)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        suffix = path.suffix.lower()
        if suffix in {".txt", ".md", ".rtf"}:
            return path.read_text(encoding="utf-8", errors="ignore")

        if suffix == ".pdf":
            reader = PdfReader(str(path))
            pages = []
            for page in reader.pages:
                page_text = page.extract_text() or ""
                pages.append(page_text)
            return "\n".join(pages)

        if suffix == ".docx":
            doc = DocxDocument(str(path))
            return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())

        raise ValueError(f"Unsupported file type: {suffix}")

    def _extract_text_from_bytes(self, raw_bytes: bytes, filename: str) -> str:
        """Helper to parse text from bytes using the filename extension."""
        if not raw_bytes:
            return ""

        suffix = Path(filename).suffix.lower()
        if suffix in {".txt", ".md", ".rtf"}:
            return raw_bytes.decode("utf-8", errors="ignore")

        if suffix == ".pdf":
            reader = PdfReader(BytesIO(raw_bytes))
            pages = []
            for page in reader.pages:
                page_text = page.extract_text() or ""
                pages.append(page_text)
            return "\n".join(pages)

        if suffix == ".docx":
            doc = DocxDocument(BytesIO(raw_bytes))
            return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())

        raise ValueError(f"Unsupported file type: {suffix}")

    def process_document(
        self,
        file_content: Union[bytes, str],
        filename: str,
        user_id: int,
        db: Session,
        create_chunks: bool = False,
    ):
        """Extract text, store the document, and optionally create chunk records."""
        if isinstance(file_content, str):
            text_content = file_content
            raw_bytes = file_content.encode("utf-8")
        else:
            text_content = self._extract_text_from_bytes(file_content, filename)
            raw_bytes = file_content

        document = UserDocument(
            filename=filename,
            file_path=f"uploads/{filename}",
            file_size=len(raw_bytes),
            user_id=user_id,
            content=text_content,
            content_preview=text_content[:500],
            processed=True,
        )
        db.add(document)
        db.commit()
        db.refresh(document)

        if create_chunks:
            chunks = self.chunker.process_document_chunks(document.id, text_content, db)
            return document, text_content, chunks

        return document, text_content

    def get_user_documents(self, user_id: int, db: Session) -> List[UserDocument]:
        """Return all documents belonging to a user."""
        return db.query(UserDocument).filter(UserDocument.user_id == user_id).all()

